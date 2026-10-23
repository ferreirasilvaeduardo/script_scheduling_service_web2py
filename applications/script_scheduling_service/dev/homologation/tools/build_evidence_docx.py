# -*- coding: utf-8 -*-
"""Gera o documento de evidencias da homologacao (DEV real) a partir do results.json do capture_evidence.js.

Uso (da pasta tools/, Python 3 com python-docx):
    SSS_EVIDENCE_DATE=AAAAMMDD python3 build_evidence_docx.py [pasta_evidencias] [saida.docx]
Rodadas conhecidas em RODADAS (20261022: CR-016; 20261023: CR-017). Padrao: a rodada de 20261023,
../evidence_<data> -> ../../../docs/homologation/SSS_Homologation_Evidence_<AAAA-MM-DD>.docx
"""
import io
import json
import os
import re
import calendar
import struct
import sys
import time

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
HOMOLOG = os.path.dirname(AQUI)
APP = os.path.dirname(os.path.dirname(HOMOLOG))
RUN_DATE = os.environ.get("SSS_EVIDENCE_DATE", "20261023")
DATE_ISO = "%s-%s-%s" % (RUN_DATE[:4], RUN_DATE[4:6], RUN_DATE[6:])
DATE_BR = "%s/%s/%s" % (RUN_DATE[6:], RUN_DATE[4:6], RUN_DATE[:4])
STAMP = "HOMOLOG_" + RUN_DATE
EVID = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HOMOLOG, "evidence_" + RUN_DATE)
SAIDA = sys.argv[2] if len(sys.argv) > 2 else os.path.join(APP, "docs", "homologation",
                                                          "SSS_Homologation_Evidence_%s.docx" % DATE_ISO)
ROXO = RGBColor(0x4A, 0x0A, 0x64)
VERDE = RGBColor(0x1B, 0x6E, 0x3A)
VERMELHO = RGBColor(0xB0, 0x1C, 0x1C)
AMBAR = RGBColor(0x8A, 0x5A, 0x00)

ORDEM = ["CT-01", "CT-02", "CT-03", "MASSA", "CT-04", "CT-05", "CT-06", "CT-07", "CT-08", "CT-09", "CT-10",
         "CT-11", "CT-12", "CT-13", "CT-14", "CT-15", "CT-16", "CT-17", "CT-18", "CT-19", "CT-20", "CT-21",
         "CT-22", "CT-23", "CT-24", "CT-25", "CT-26", "ENCERRAMENTO"]

# Adaptacao de cada caso ao DEV real (o manual de 18/10 foi escrito sobre a massa simulada do FAKE).
ADAPTACAO = {
    "CT-01": u"Igual ao manual. A troca válida de senha do admin usou a senha escolhida pelo responsável pela "
             u"homologação; os demais usuários seguem com a senha inicial documentada (DEPLOYMENT.md §5.3).",
    "CT-03": u"HOMOLOGACAO → GFPREPROD; o perfil OPERATOR_HML do DEV tem CLONE6, CLONE7 e GFPREPROD.",
    "CT-04": u"carga_diaria → contar_linhas_arquivo; o histórico mostra as execuções reais do dia.",
    "CT-05": u"carga_diaria em PRODUCAO → gerar_hash_arquivo em CLONE2. Os 14 utilitários do seed são liberados "
             u"aos três perfis operacionais, então não há script fora do perfil para mostrar na etapa 1.",
    "CT-06": u"O ticket rodando é uma compactação real de 3 GB de texto (nível 9, cerca de 4 minutos).",
    "CT-07": u"Os tickets #11 e #30 do FAKE → ticket longo rodando e um ticket que deu erro de verdade "
             u"(arquivo inexistente, código de retorno 2). A exibição segura do log usa um caminho com "
             u"<script> e <b>, ecoado pelo próprio script.",
    "CT-08": u"O ticket esperando foi criado com Agendar para daqui a 2 horas (o agente real pega os imediatos em "
             u"segundos).",
    "CT-09": u"Ticket #19 do FAKE → ticket real do op.prod (gerar_hash_arquivo em CLONE2); o vínculo obrigatório "
             u"temporário foi o termo_busca, desfeito ao final.",
    "CT-11": u"HOMOLOG2 → ambiente %s_<hhmm>, desativado ao final; a confirmação de desativação "
             u"foi aberta no GFPREPROD e cancelada." % STAMP,
    "CT-12": u"CONSULTA → READ_ONLY; OPERADOR_HML → OPERATOR_HML; DESENV (inativo) → CLONE7; VER_AUDITORIA → "
             u"VIEW_AUDIT. A revisão foi cancelada (matriz inalterada).",
    "CT-13": u"carga_diaria → contar_linhas_arquivo; HOMOLOGACAO → GFPREPROD.",
    "CT-14": u"O usuário bloqueado foi produzido de verdade: 5 tentativas erradas do consulta, depois "
             u"desbloqueado pela tela. A redefinição válida do op.prod usou a própria senha inicial.",
    "CT-15": u"data_ref/documento_cliente → parâmetro sigiloso documento_cliente_hml criado pela tela "
             u"(regex ^[0-9]{11}$), sem vínculo com scripts.",
    "CT-16": u"carga_diaria → contar_linhas_arquivo; o caminho original foi restaurado no fim. Pela regra "
             u"atual (CR-011), o arquivo é conferido pelo agente na execução, não pelo nó web.",
    "CT-17": u"Regra nova %s_DIARIA_<hhmm>, diária com a primeira execução em poucos minutos; o "
             u"agente real criou o ticket no horário (CT-18)." % STAMP,
    "CT-18": u"A regra pausada do FAKE não existe no DEV; a agenda mostra a regra do CT-17.",
    "CT-19": u"Códigos atuais da auditoria: TICKET_CANCEL e AUDIT_EXPORT (antes TICKET_CANCELAR e AUD_EXPORTAR).",
    "CT-20": u"Um único agente real (srv-linux-local-01); os agentes agt-host01/02 eram simulados pelo FAKE.",
    "CT-21": u"Ticket fora do escopo do op.hml: um ticket real de CLONE2.",
    "CT-23": u"PRODUCAO → ambiente criado no CT-11; ticket #17 → ticket real do op.prod.",
    "CT-24": u"Responsável inativo e link simbólico seguem a regra atual (ver Observações).",
    "CT-25": u"SHA-256 retirado pelo CR-011; a matriz do OPERATOR_PROD voltou ao original no fim.",
}

# Ressalvas (o caso passou, com observacao relevante para o responsavel).
_RESSALVA_1022 = {
    "CT-05": u"Destaque de produção fixo no nome PRODUCAO (ticket_service.PRODUCTION_ENVIRONMENT): não aparece "
             u"para nenhum ambiente do DEV.",
    "CT-20": u"Estados Atrasado/Em risco/Offline e quadros de atenção não reproduzíveis com um único agente "
             u"saudável.",
    "CT-24": u"Cenário de responsável inativo não reproduzível no DEV sem mudar permissões (coberto pelos testes "
             u"automatizados).",
}

_DEFEITOS_1022 = [
    (u"Linha de situação ao vivo vazia em 4 telas",
     u"views/tickets/index.html, views/tickets/detail.html, views/audit/agents.html, views/admin_profiles/matrix.html",
     u"O template mostrava x-text=\"state\", mas os componentes Alpine expõem statusText. Desde o CR-016a não "
     u"apareciam: \"Atualizada às …\" da lista de tickets, \"Atualizado às …\" do painel de agentes, a situação do "
     u"visualizador de log (\"Acompanhando a execução\", \"Atualização pausada\") e a contagem de alterações "
     u"pendentes da matriz.",
     u"Corrigido; conferido no CT-06, CT-12 e CT-20."),
    (u"Duração típica \"cerca de menos de 1 minuto\"",
     u"modules/ticket_service.py (histórico do script)",
     u"O prefixo \"cerca de\" era somado ao texto \"menos de 1 minuto\".",
     u"Corrigido (abaixo de 1 minuto: \"menos de 1 minuto\"); teste novo em test_catalog_service."),
    (u"Caminho longo estoura a janela de confirmação do disparo",
     u"static/css/sss.css (.sss-summary-answers)",
     u"As respostas do resumo (caminhos de arquivo sem espaço) não quebravam linha e passavam da borda do modal.",
     u"Corrigido (quebra em qualquer ponto); conferido na figura de confirmação do CT-05."),
]

_OBSERVACOES_1022 = [
    u"Destaque de ambiente de produção: só o ambiente chamado PRODUCAO recebe o selo vermelho e a confirmação "
    u"\"AMBIENTE DE PRODUÇÃO\" (ticket_service.PRODUCTION_ENVIRONMENT). No DEV nenhum ambiente tem esse nome; "
    u"recomenda-se torná-lo configurável antes da implantação em produção.",
    u"Cancelamento no agente real: o log não ganha a linha \"processo encerrado por cancelamento (SIGTERM)\", "
    u"que só existia na simulação do FAKE; o motivo aparece no Detalhe técnico (CANCELLED / CANCEL_USER).",
    u"Scripts PYTHON3 rodam com -I e saída em buffer: o log de um script Python aparece quando o buffer descarrega "
    u"(no fim, ou nunca, se o processo for cancelado). Scripts SH/BASH escrevem na hora.",
    u"Sufixo do id da mensagem de erro dos campos: components/form_field.html usa \"-erro\" e os formulários próprios "
    u"(disparo, regra) usam \"-error\". Sem efeito para o usuário; resquício do CR-016, cuja tradução foi encerrada.",
]


# Rodada de 23/10/2026 (CR-017): o que mudou nas telas e a massa de um banco DEV recriado.
_ADAPTACAO_1023 = {
    "CT-02": u"Painel sem o cartão de disparo e sem o menu repetido (CR-017b): o disparo fica no item de menu "
             u"Rodar um script.",
    "CT-04": u"Rotinas disponíveis passou a se chamar Rodar um script (CR-017b), com o botão Agendar para quem tem "
             u"Regras de Agendamento (CR-017c).",
    "CT-05": u"carga_diaria em PRODUCAO → gerar_hash_arquivo em CLONE2. A escolha do script é feita nos cartões de "
             u"Rodar um script; o comando aparece só na confirmação (sem a pré-visualização separada, CR-017b). "
             u"Todos os ambientes do DEV estão marcados como produção no cadastro (CR-017a), então o alerta e o botão "
             u"vermelho aparecem.",
    "CT-26": u"Caso novo do CR-017: o perfil OPERATOR_HML perdeu VIEW_LOGS por alguns segundos (para mostrar o item "
             u"Tickets de quem só dispara) e o limpar_linhas_vazias ficou homologado em um ambiente e depois em "
             u"nenhum (disparo e regra direto ao formulário; quadro Quem pode rodar). Tudo foi desfeito no próprio "
             u"caso.",
}

_OBSERVACOES_1023 = [
    u"Banco do DEV recriado no CR-017a (rollback, implantação e validação dos scripts de docs/sql/): os "
    u"registros da rodada de 22/10 não existem mais, e a massa desta rodada foi criada do zero pelas telas.",
    u"Ambiente de produção agora é uma marcação do cadastro (tbsss_environments.is_production, CR-017a). O seed "
    u"marca os 5 ambientes do DEV como produção; por isso o alerta aparece em todo disparo.",
    u"Cancelamento no agente real: o log não ganha a linha \"processo encerrado por cancelamento (SIGTERM)\", "
    u"que só existia na simulação do FAKE; o motivo aparece no Detalhe técnico (CANCELLED / CANCEL_USER).",
    u"Scripts PYTHON3 rodam com -I e saída em buffer: o log de um script Python aparece quando o buffer descarrega "
    u"(no fim, ou nunca, se o processo for cancelado). Scripts SH/BASH escrevem na hora.",
    u"Três verificações da primeira passada (menu com Administração, alerta de produção e coluna Produção) "
    u"procuravam o texto como escrito, mas a tela o exibe em maiúsculas por CSS. O roteiro passou a comparar sem "
    u"diferenciar maiúsculas, e o CT-02, o CT-05 e o CT-11 foram executados de novo; por isso há um segundo "
    u"ambiente %s_<hhmm> (também inativo) e um segundo ticket do CT-05." % STAMP,
]

RODADAS = {
    "20261022": {"commit": u"07e424f (encerramento do CR-016) mais as correções desta homologação, no mesmo commit "
                           u"deste documento",
                 "adaptacao": {}, "ressalva": _RESSALVA_1022, "defeitos": _DEFEITOS_1022,
                 "observacoes": _OBSERVACOES_1022},
    "20261023": {"commit": u"dedf93a (CR-017 concluído: 017a a 017d) mais o roteiro ajustado, no mesmo commit "
                           u"deste documento",
                 "adaptacao": _ADAPTACAO_1023,
                 "ressalva": dict((k, v) for k, v in _RESSALVA_1022.items() if k != "CT-05"),
                 "defeitos": [], "observacoes": _OBSERVACOES_1023},
}
RODADA = RODADAS[RUN_DATE]
ADAPTACAO.update(RODADA["adaptacao"])
RESSALVA = RODADA["ressalva"]
DEFEITOS = RODADA["defeitos"]
OBSERVACOES = RODADA["observacoes"]

def local(iso):
    """ISO em UTC (registro do roteiro) -> 'dd/mm/aaaa hh:mm:ss' no horario desta maquina (o mesmo das telas)."""
    utc = calendar.timegm(time.strptime(iso[:19], "%Y-%m-%dT%H:%M:%S"))
    return time.strftime("%d/%m/%Y %H:%M:%S", time.localtime(utc))


def png_size(path):
    with open(path, "rb") as f:
        head = f.read(24)
    return struct.unpack(">II", head[16:24])


def shade(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def set_col_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return p


def kv_table(doc, rows, widths=(4.2, 12.8)):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in rows:
        r = t.add_row().cells
        cell_text(r[0], k, bold=True)
        shade(r[0], "EFE7F3")
        cell_text(r[1], v)
    set_col_widths(t, widths)
    doc.add_paragraph()
    return t


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
    run._r.append(fld)
    doc.add_paragraph(u"Se o sumário aparecer vazio, clique com o botão direito sobre ele e escolha "
                      u"\"Atualizar campo\".").runs[0].italic = True


def resultado_ct(ct_id, r):
    ok = all(c["ok"] for c in r["checks"])
    if not ok:
        return u"Não OK", VERMELHO
    if ct_id in RESSALVA:
        return u"OK com ressalva", AMBAR
    return u"OK", VERDE


def passos_manual():
    """Objetivo e usuarios de cada CT, do texto do manual (all_cts_steps.txt)."""
    txt = io.open(os.path.join(HOMOLOG, "all_cts_steps.txt"), encoding="utf-8").read()
    out = {}
    partes = re.split(r"^=== (CT-\d+\.?) ===$", txt, flags=re.M)
    for i in range(1, len(partes), 2):
        cid, corpo = partes[i], partes[i + 1]
        if cid.endswith("."):
            continue
        linhas = [l.strip() for l in corpo.splitlines() if l.strip()]
        obj = linhas[linhas.index("Objetivo") + 1] if "Objetivo" in linhas else ""
        usr = linhas[linhas.index(u"Usuário(s)") + 1] if u"Usuário(s)" in linhas else ""
        out[cid] = (obj, usr)
    return out


def main():
    res = json.load(io.open(os.path.join(EVID, "results.json"), encoding="utf-8"))
    state = json.load(io.open(os.path.join(EVID, "state.json"), encoding="utf-8"))
    manual = passos_manual()
    ids = [c for c in ORDEM if c in res]
    inicio = min(res[c]["started_at"] for c in ids)
    fim = max(res[c].get("finished_at", res[c]["started_at"]) for c in ids)
    total_chk = sum(len(res[c]["checks"]) for c in ids)
    total_ok = sum(1 for c in ids for k in res[c]["checks"] if k["ok"])
    total_fig = sum(len(res[c]["figures"]) for c in ids)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(2.0))
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        doc.styles[name].font.color.rgb = ROXO

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run(u"Script Scheduling Service (SSS)")
    r.font.size = Pt(14)
    r.font.color.rgb = ROXO
    doc.add_heading(u"Evidências de Homologação — DEV real (Oracle e agente)", 0)
    doc.add_paragraph(
        u"Execução real dos casos de teste do Manual de Homologação (UI Freeze, versão 2.5) sobre o ambiente DEV: "
        u"banco Oracle, servidor web na porta 18000 e o agente srv-linux-local-01 executando os scripts de verdade. "
        u"Cada caso foi executado no navegador por um roteiro automatizado, que conferiu na própria página o "
        u"resultado esperado e capturou as telas.")
    kv_table(doc, [
        (u"Data da execução", u"%s a %s (horário de Brasília, o mesmo das telas)" % (local(inicio), local(fim)[11:])),
        (u"Ambiente", u"DEV — web2py (Python 3.6.15) em http://127.0.0.1:18000, banco Oracle 19c (driver thick)"),
        (u"Agente", u"srv-linux-local-01 (Linux), ambientes CLONE2, CLONE1, CLONE6, CLONE7 e GFPREPROD; "
                    u"interpretadores PYTHON3, BASH e SH"),
        (u"Base", u"Manual de Homologação — UI Freeze 2.5 (17/10/2026), adaptado aos dados do DEV"),
        (u"Commit de referência", RODADA["commit"]),
        (u"Execução", u"Roteiro dev/homologation/tools/capture_evidence.js (Chrome headless), %d verificações, "
                      u"%d figuras" % (total_chk, total_fig)),
        (u"Resultado geral", u"%d de %d verificações OK" % (total_ok, total_chk)),
        (u"Responsável pela homologação", u""),
    ])

    doc.add_heading(u"Sumário", 1)
    add_toc(doc)
    doc.add_page_break()

    # 1. Resumo
    doc.add_heading(u"1. Resumo da execução", 1)
    tb = doc.add_table(rows=1, cols=5)
    tb.style = "Table Grid"
    for i, h in enumerate((u"Caso", u"Título", u"Verificações", u"Figuras", u"Resultado")):
        cell_text(tb.rows[0].cells[i], h, bold=True)
        shade(tb.rows[0].cells[i], "4A0A64")
        tb.rows[0].cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for c in ids:
        r = res[c]
        ok = sum(1 for k in r["checks"] if k["ok"])
        txt, cor = resultado_ct(c, r)
        row = tb.add_row().cells
        cell_text(row[0], c if c.startswith("CT") else u"—")
        cell_text(row[1], r["title"])
        cell_text(row[2], u"%d/%d" % (ok, len(r["checks"])))
        cell_text(row[3], str(len(r["figures"])))
        cell_text(row[4], txt, bold=True, color=cor)
    set_col_widths(tb, (1.6, 9.2, 2.2, 1.6, 2.4))
    doc.add_paragraph()

    # 2. Preparacao
    doc.add_heading(u"2. Preparação do ambiente", 1)
    doc.add_heading(u"2.1 Servidor web e agente", 2)
    for txt in (
            u"Servidor web do DEV iniciado com bin/start_local.sh 18000 (Python 3.6.15), ligado ao Oracle do DEV.",
            u"Agente iniciado com bin/sss-scheduler.sh e uma cópia do instance/config.ini do DEV com o mesmo agent_id "
            u"(srv-linux-local-01), os mesmos ambientes e as mesmas pastas, acrescentando apenas "
            u"interpreters = PYTHON3=/usr/bin/python3,BASH=/bin/bash,SH=/bin/sh (o padrão é vazio, e sem "
            u"interpretadores o agente não pega tickets).",
            u"Scripts do seed: o cadastro aponta para /opt/automacoes/scripts/utilitarios/, instalado nesta máquina "
            u"como link simbólico para dev/scripts/utilitarios/ do repositório (aceito pelo CR-011).",
            u"Arquivos de entrada reais nas pastas de entrada de cada ambiente e script (CR-014): clientes, notas e "
            u"extrato em texto e um lote de 3 GB de texto para a compactação longa."):
        doc.add_paragraph(txt, style="List Bullet")
    doc.add_heading(u"2.2 Usuários e senhas", 2)
    doc.add_paragraph(
        u"Usuários do seed do DEV: admin (ADMINISTRATOR), op.prod (OPERATOR_PROD), op.hml (OPERATOR_HML) e "
        u"consulta (READ_ONLY), com a senha inicial documentada no DEPLOYMENT.md §5.3. No CT-01 a senha do admin foi "
        u"trocada pela tela para a senha escolhida pelo responsável pela homologação. Nenhuma senha aparece neste "
        u"documento, nas telas capturadas ou no roteiro.")
    doc.add_heading(u"2.3 Massa criada pelas telas", 2)
    m = res.get("MASSA")
    if m:
        doc.add_paragraph(u"Antes do CT-04, cinco tickets reais foram criados pelo assistente de disparo e executados "
                          u"pelo agente:")
        for k in m["checks"]:
            if k["desc"].startswith(u"Ticket"):
                doc.add_paragraph(k["desc"] + (u" — " + k["detail"] if k.get("detail") else u""), style="List Bullet")
    doc.add_heading(u"2.4 Registros mantidos como evidência", 2)
    doc.add_paragraph(
        u"Os registros criados ficam no banco do DEV, identificados pela marca %s: tickets e "
        u"execuções, o ambiente %s_<hhmm> (inativo), o parâmetro sigiloso documento_cliente_hml, o "
        u"parâmetro ct24_valor_<hhmm> e as regras %s_DIARIA_<hhmm> (pausadas no encerramento). "
        % (STAMP, STAMP, STAMP) +
        u"Alterações de cadastro feitas para testar (caminho de script, vínculo de parâmetro, permissão "
        u"MANAGE_USERS do OPERATOR_PROD) foram desfeitas pela tela no próprio caso. Tudo está na trilha de "
        u"auditoria.")

    # 3. Defeitos e observacoes
    doc.add_heading(u"3. Defeitos encontrados e corrigidos", 1)
    td = doc.add_table(rows=1, cols=4)
    td.style = "Table Grid"
    for i, h in enumerate((u"Defeito", u"Onde", u"Causa", u"Situação")):
        cell_text(td.rows[0].cells[i], h, bold=True)
        shade(td.rows[0].cells[i], "EFE7F3")
    for d in DEFEITOS:
        row = td.add_row().cells
        for i, v in enumerate(d):
            cell_text(row[i], v)
    set_col_widths(td, (3.6, 4.0, 6.0, 3.4))
    doc.add_paragraph()
    doc.add_heading(u"4. Observações e diferenças em relação ao manual", 1)
    for o in OBSERVACOES:
        doc.add_paragraph(o, style="List Bullet")

    # 5. Casos
    doc.add_page_break()
    doc.add_heading(u"5. Casos de teste executados", 1)
    fig_n = 0
    for c in ids:
        r = res[c]
        doc.add_heading((u"%s — %s" % (c, r["title"])) if c.startswith("CT") else r["title"], 2)
        obj, usr = manual.get(c, (u"", u""))
        linhas = []
        if obj:
            linhas.append((u"Objetivo", obj))
        if usr:
            linhas.append((u"Usuário(s)", usr))
        if c in ADAPTACAO:
            linhas.append((u"Adaptação ao DEV", ADAPTACAO[c]))
        linhas.append((u"Executado em", local(r["started_at"])))
        if linhas:
            kv_table(doc, linhas)
        doc.add_heading(u"Verificações executadas", 3)
        tv = doc.add_table(rows=1, cols=4)
        tv.style = "Table Grid"
        for i, h in enumerate((u"#", u"Verificação (resultado esperado)", u"Obtido", u"Resultado")):
            cell_text(tv.rows[0].cells[i], h, bold=True)
            shade(tv.rows[0].cells[i], "EFE7F3")
        for n, k in enumerate(r["checks"], 1):
            row = tv.add_row().cells
            cell_text(row[0], str(n))
            cell_text(row[1], k["desc"])
            cell_text(row[2], k.get("detail", u""), size=8)
            cell_text(row[3], u"OK" if k["ok"] else u"Não OK", bold=True, color=VERDE if k["ok"] else VERMELHO)
        set_col_widths(tv, (0.8, 8.0, 6.2, 2.0))
        if r["notes"]:
            doc.add_heading(u"Observações", 3)
            for n in r["notes"]:
                doc.add_paragraph(n, style="List Bullet")
        if r["figures"]:
            doc.add_heading(u"Telas", 3)
        for f in r["figures"]:
            fig_n += 1
            path = os.path.join(EVID, f["file"])
            w, h = png_size(path)
            largura = 16.5
            altura = largura * h / float(w)
            if altura > 21.5:
                largura = 21.5 * w / float(h)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(path, width=Cm(largura))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = cap.add_run(u"Figura %d — %s" % (fig_n, f["caption"]))
            rr.italic = True
            rr.font.size = Pt(9)
            u = cap.add_run(u"  (%s)" % f["url"])
            u.font.size = Pt(8)
            u.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        doc.add_heading(u"Registro", 3)
        txt, cor = resultado_ct(c, r)
        reg = [(u"Resultado", txt)]
        if c in RESSALVA:
            reg.append((u"Ressalva", RESSALVA[c]))
        reg.append((u"Executado por", u"Roteiro automatizado capture_evidence.js, em %s" % DATE_BR))
        reg.append((u"Conferido por / data", u""))
        tr = kv_table(doc, reg)
        tr.rows[0].cells[1].paragraphs[0].runs[0].bold = True
        tr.rows[0].cells[1].paragraphs[0].runs[0].font.color.rgb = cor

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    doc.save(SAIDA)
    resumo_md(res, ids, inicio, fim, total_ok, total_chk, fig_n)
    print("gerado:", SAIDA, "| casos:", len(ids), "| verificacoes:", total_ok, "/", total_chk, "| figuras:", fig_n)


def resumo_md(res, ids, inicio, fim, total_ok, total_chk, total_fig):
    """Resumo em Markdown ao lado do .docx (como o registro do UI Freeze de 18/10)."""
    nao_ok = [(c, k) for c in ids for k in res[c]["checks"] if not k["ok"]]
    linhas = [
        u"# Homologação no DEV real: Registro de Evidências", u"",
        u"**Data:** %s a %s (horário de Brasília)" % (local(inicio), local(fim)[11:]), u"",
        u"**Ambiente:** DEV, banco Oracle, web2py em `127.0.0.1:18000`, agente `srv-linux-local-01`",
        u"**Resultado geral:** %d de %d verificações OK; %d figuras." % (total_ok, total_chk, total_fig), u"",
        u"Documento completo, com as telas: `%s` (mesma pasta). "
        u"Gerado por `dev/homologation/tools/capture_evidence.js` e `build_evidence_docx.py`; o resultado bruto "
        u"de cada verificação está em `dev/homologation/evidence_%s/results.json`. Nenhuma senha foi "
        u"registrada." % (os.path.basename(SAIDA), RUN_DATE), u"",
        u"## Resultado por caso", u"",
        u"| Caso | Título | Verificações | Figuras | Resultado |", u"|---|---|---|---|---|",
    ]
    for c in ids:
        r = res[c]
        txt, _cor = resultado_ct(c, r)
        linhas.append(u"| %s | %s | %d/%d | %d | %s |" % (c if c.startswith("CT") else u"—", r["title"],
                      sum(1 for k in r["checks"] if k["ok"]), len(r["checks"]), len(r["figures"]), txt))
    linhas += [u"", u"## Defeitos encontrados e corrigidos", u""]
    for d in DEFEITOS:
        linhas.append(u"- **%s** (`%s`): %s %s" % d)
    if not DEFEITOS:
        linhas.append(u"Nenhum.")
    linhas += [u"", u"## Observações e diferenças em relação ao manual", u""]
    linhas += [u"- " + o for o in OBSERVACOES]
    linhas += [u"", u"## Ressalvas", u""]
    linhas += [u"- **%s:** %s" % (c, t) for c, t in sorted(RESSALVA.items())]
    linhas += [u"", u"## Verificações que não passaram", u""]
    linhas += ([u"- %s: %s (%s)" % (c, k["desc"], k.get("detail", u"")) for c, k in nao_ok] or [u"Nenhuma."])
    md = os.path.join(os.path.dirname(SAIDA), "SSS_HOMOLOGATION_EVIDENCE_%s.md" % DATE_ISO)
    io.open(md, "w", encoding="utf-8").write(u"\n".join(linhas) + u"\n")
    print("gerado:", md)


if __name__ == "__main__":
    main()
